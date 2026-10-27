"""
Testes da Etapa E2 — Importação.
Valida parser de planilhas, detecção de cabeçalho, extração de itens e notas históricas.
RN19: Valores digitados são respeitados como estão.
CA12/CA13: Lê totais digitados à mão e valores em texto.
"""
import pytest
from decimal import Decimal
from datetime import date
from unittest.mock import MagicMock, patch
from importacao.parsers import (
    detectar_cabecalho,
    detectar_colunas_notas,
    extrair_itens_planilha,
    parse_valor,
)


class TestParseValor:
    """Testes da função parse_valor."""

    def test_none_retorna_none(self):
        assert parse_valor(None) is None

    def test_string_vazia_retorna_none(self):
        assert parse_valor("") is None
        assert parse_valor("   ") is None

    def test_inteiro(self):
        assert parse_valor(100) == Decimal("100")

    def test_float(self):
        assert parse_valor(10.5) == Decimal("10.5")

    def test_string_numero_simples(self):
        assert parse_valor("1234") == Decimal("1234")

    def test_string_formato_brasileiro(self):
        assert parse_valor("1.234,56") == Decimal("1234.56")

    def test_string_com_cifrao_e_espacos(self):
        assert parse_valor("R$ 1.234,56") == Decimal("1234.56")

    def test_ca13_valor_em_texto(self):
        """CA13: Lê valores gravados como texto ('R$16.752,00')."""
        assert parse_valor("R$16.752,00") == Decimal("16752.00")

    def test_string_invalida_retorna_none(self):
        assert parse_valor("abc") is None
        assert parse_valor("não é número") is None


class TestDetectarCabecalho:
    """Testes da detecção automática de cabeçalho."""

    def _make_ws(self, rows_data):
        """Cria um mock de worksheet com dados simulados."""
        ws = MagicMock()
        ws.max_row = len(rows_data)

        def get_row(idx):
            if idx <= len(rows_data):
                cells = []
                for col_idx, val in enumerate(rows_data[idx - 1]):
                    cell = MagicMock()
                    cell.value = val
                    cell.column_letter = chr(ord("A") + col_idx)
                    cells.append(cell)
                return cells
            return []

        ws.__getitem__ = lambda self, key: get_row(key)
        ws.iter_rows = lambda **kwargs: [
            [MagicMock(value=v, column_letter=chr(ord("A") + i)) for i, v in enumerate(row)]
            for row in rows_data
        ]
        return ws

    def test_detecta_cabecalho_padrao(self):
        """Detecta cabeçalho com ITEM, DESCRIÇÃO, VLR UNT, QNT LICITADA."""
        rows = [
            [None, None, None],
            ["ITEM", "DESCRIÇÃO", "MARCA", "UND", "VLR UNT", "QNT LICITADA"],
            [1, "Parafuso", "X", "UN", 0.50, 100],
        ]
        ws = self._make_ws(rows)
        linha, mapeamento = detectar_cabecalho(ws)
        assert linha == 2
        assert "descricao" in mapeamento
        assert "item" in mapeamento

    def test_detecta_cabecalho_objeto(self):
        """Detecta 'OBJETO' como descrição."""
        rows = [
            ["ACOMPANHAMENTO CONTRATOS"],
            ["ITEM", "OBJETO", "VALOR", "QUANTIDADE"],
            [1, "Material", 10.00, 50],
        ]
        ws = self._make_ws(rows)
        linha, mapeamento = detectar_cabecalho(ws)
        assert "descricao" in mapeamento

    def test_retorna_vazio_sem_cabecalho(self):
        """Retorna mapeamento vazio se não encontrar cabeçalho."""
        rows = [
            [None, None, None],
            [None, None, None],
        ]
        ws = self._make_ws(rows)
        linha, mapeamento = detectar_cabecalho(ws)
        assert mapeamento == {} or "descricao" not in mapeamento


class TestDetectarColunasNotas:
    """Testes da detecção de colunas de notas fiscais históricas."""

    def test_detecta_nfe_no_cabecalho(self):
        """Detecta colunas com padrão 'NFE 266'."""
        ws = MagicMock()
        header_cells = []
        for col_idx, val in enumerate(["ITEM", "OBJETO", "NFE 266", "NFE 250"]):
            cell = MagicMock()
            cell.value = val
            cell.column_letter = chr(ord("A") + col_idx)
            header_cells.append(cell)
        ws.__getitem__ = lambda self, key: header_cells if key == 1 else []

        notas = detectar_colunas_notas(ws, 1)
        assert len(notas) == 2
        numeros = {n["numero_nfe"] for n in notas}
        assert "266" in numeros
        assert "250" in numeros

    def test_sem_notas_retorna_vazio(self):
        """Retorna lista vazia quando não há colunas de NFE."""
        ws = MagicMock()
        header_cells = []
        for col_idx, val in enumerate(["ITEM", "OBJETO", "VALOR"]):
            cell = MagicMock()
            cell.value = val
            cell.column_letter = chr(ord("A") + col_idx)
            header_cells.append(cell)
        ws.__getitem__ = lambda self, key: header_cells if key == 1 else []

        notas = detectar_colunas_notas(ws, 1)
        assert notas == []


@pytest.mark.django_db
class TestExtrairItensPlanilha:
    """Testes de extração de itens usando planilha real de fixture."""

    def test_extrai_itens_de_planilha_real(self):
        """Extrai itens da planilha ECO MAIS com sucesso."""
        filepath = "planilhas em uso/2022/CONTROLES CONTRATO/ECO MAIS AR CONDICIONADOS LTDA.xlsx"
        # Usa a sheet '010607-2021' que tem dados reais
        import openpyxl

        wb = openpyxl.load_workbook(filepath, data_only=True)
        ws = wb["010607-2021"]
        linha_cabecalho, mapeamento = detectar_cabecalho(ws)
        wb.close()

        assert "descricao" in mapeamento, f"Mapeamento encontrado: {mapeamento}"

        itens, avisos = extrair_itens_planilha(filepath, mapeamento, linha_cabecalho + 1)
        # A planilha tem pelo menos alguns itens com descrição
        assert len(itens) > 0, f"Nenhum item extraído. Avisos: {avisos}"

    def test_rn19_valores_digitados_respeitados(self):
        """RN19: Valores unitários digitados são lidos exatamente."""
        resultado = parse_valor("R$1485,00")
        assert resultado == Decimal("1485.00")

        resultado = parse_valor(1485)
        assert resultado == Decimal("1485")